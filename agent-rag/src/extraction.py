"""Binary document text extraction.

Supports PDF (via pypdf) and Word docx (via python-docx).
Both libraries are optional — if not installed the function falls back to
UTF-8 decode with errors='ignore', which is the original behaviour for
plain-text formats.

Position-aware extraction:
  extract_with_position() returns a list of PageSegment dicts so callers
  can attach page/line metadata to each chunk for source traceability.
"""
import io
from typing import TypedDict


class PageSegment(TypedDict):
    """One page (or logical section) extracted from a document."""
    page: int          # 1-based page number
    lines: list[str]   # non-empty lines in order
    text: str          # full text of this segment (lines joined by \n)


def extract_text(content: bytes, file_type: str) -> str:
    """Return plain text from *content* bytes (flat string, no position info)."""
    ft = file_type.lower()
    if ft == "pdf":
        return _extract_pdf(content)
    if ft in ("docx", "doc"):
        return _extract_docx(content)
    if ft in ("xlsx", "xls"):
        return _extract_xlsx(content)
    return content.decode("utf-8", errors="ignore")


def extract_with_position(content: bytes, file_type: str) -> list[PageSegment]:
    """Return structured per-page segments with line information.

    Falls back to a single PageSegment when the format has no page concept
    (plain text, docx, xlsx).
    """
    ft = file_type.lower()
    if ft == "pdf":
        return _extract_pdf_pages(content)
    # For non-PDF formats produce a single synthetic segment so the pipeline
    # can use the same code path regardless of file type.
    text = extract_text(content, file_type)
    lines = [l for l in text.splitlines() if l.strip()]
    return [PageSegment(page=1, lines=lines, text=text)]


# ── PDF ──────────────────────────────────────────────────────────────────────

def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return content.decode("utf-8", errors="ignore")

    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def _extract_pdf_pages(content: bytes) -> list[PageSegment]:
    try:
        from pypdf import PdfReader
    except ImportError:
        text = content.decode("utf-8", errors="ignore")
        lines = [l for l in text.splitlines() if l.strip()]
        return [PageSegment(page=1, lines=lines, text=text)]

    reader = PdfReader(io.BytesIO(content))
    segments: list[PageSegment] = []
    for page_no, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        lines = [l for l in raw.splitlines() if l.strip()]
        if lines:
            segments.append(PageSegment(page=page_no, lines=lines, text="\n".join(lines)))
    return segments


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        return content.decode("utf-8", errors="ignore")

    doc = Document(io.BytesIO(content))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    """Convert xlsx to TSV-like text so the table splitter can handle it."""
    try:
        import openpyxl
    except ImportError:
        return content.decode("utf-8", errors="ignore")

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(v) for v in row if v is not None)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(parts)
